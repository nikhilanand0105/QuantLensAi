#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
QuantLens-XAI: Complete Project Report Generator
Produces Markdown (.md), Microsoft Word (.docx), and PDF (.pdf) formats.
"""

import os
import re
import subprocess
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

import build_report_content as report_data

BASE_DIR = os.path.abspath("docs/quantlens-final-report")
FIGURES_DIR = os.path.join(BASE_DIR, "figures")
os.makedirs(BASE_DIR, exist_ok=True)

MD_PATH = os.path.join(BASE_DIR, "QuantLens_XAI_Complete_Project_Report.md")
DOCX_PATH = os.path.join(BASE_DIR, "QuantLens_XAI_Complete_Project_Report.docx")
HTML_PATH = os.path.join(BASE_DIR, "QuantLens_XAI_Report_Print.html")
PDF_PATH = os.path.join(BASE_DIR, "QuantLens_XAI_Complete_Project_Report.pdf")

meta = report_data.METADATA
chapters = report_data.get_report_chapters()
appendices = report_data.get_appendices()

def sanitize_xml(s: str) -> str:
    """Removes invalid XML control characters."""
    if not isinstance(s, str):
        return str(s)
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x84\x86-\x9f]', '', s)

# ==============================================================================
# 1. GENERATE MARKDOWN REPORT (.md)
# ==============================================================================
print("[1/3] Generating Markdown Report (.md)...")
md_lines = []

md_lines.append(f"# {meta['title']}\n")
md_lines.append(f"**Course:** {meta['course']}  ")
md_lines.append(f"**Author:** {meta['student_name']} (Register Number: {meta['register_number']})  ")
md_lines.append(f"**Academic Year:** {meta['academic_year']} | **Date:** {meta['date']}  ")
md_lines.append(f"**Institution:** {meta['institution']}\n")
md_lines.append("---\n")

md_lines.append("## Table of Contents")
for ch in chapters:
    md_lines.append(f"- [{ch['title']}](#{ch['title'].lower().replace(' ', '-').replace(':', '')})")
for ap in appendices:
    md_lines.append(f"- [{ap['title']}](#{ap['title'].lower().replace(' ', '-').replace(':', '')})")
md_lines.append("\n---\n")

for ch in chapters:
    md_lines.append(f"\n## {ch['title']}\n")
    content = ch['content'].strip()
    
    # Embed figure links where relevant
    if ch['id'] == 4:
        content += f"\n\n![System Architecture Overview](figures/arch_system_overview.png)\n*Figure 4.1: QuantLens-XAI End-to-End System Architecture Overview*\n"
    elif ch['id'] == 5:
        content += f"\n\n![Hybrid Database Architecture](figures/arch_hybrid_database.png)\n*Figure 5.1: Hybrid Polyglot Persistence Architecture (SQLite 3 + MongoDB Community 7.0)*\n"
        content += f"\n\n![MongoDB Compass and CLI Inspection](figures/07_mongodb_compass_or_cli.png)\n*Figure 5.2: MongoDB Community 7.0 Collection & BSON Inspection*\n"
    elif ch['id'] == 11:
        content += f"\n\n![Strategy Builder UI](figures/01_strategy_builder.png)\n*Figure 11.1: Strategy Builder Natural Language Prompting Interface*\n"
        content += f"\n\n![7-Stage Decision Trace](figures/02_decision_trace.png)\n*Figure 11.2: Seven-Stage Observable Decision Trace and AI Risk Advisory*\n"
        content += f"\n\n![Backtest Engine Candlestick](figures/03_backtest_candlestick.png)\n*Figure 11.3: Event-Driven Backtesting Candlestick Chart, Overlays, and Trade Ledger*\n"
        content += f"\n\n![Explainable AI Lab](figures/04_xai_explainability.png)\n*Figure 11.4: Dual Model Predictions, Real SHAP Attributions, and Slide 9 Empirical Evaluation*\n"
        content += f"\n\n![Paper Broker Dashboard](figures/05_paper_broker.png)\n*Figure 11.5: Groww-Inspired Paper Broker Portfolio, Orders Ledger, and Audit Trail*\n"
        content += f"\n\n![Delta Bot Dashboard](figures/06_delta_bot_dashboard.png)\n*Figure 11.6: Legacy Delta Algorithmic Trading Bot Operations Dashboard*\n"
    
    md_lines.append(content)

md_lines.append("\n\n---\n## Appendices\n")
for ap in appendices:
    md_lines.append(f"\n### {ap['title']}\n")
    md_lines.append(ap['content'].strip())

with open(MD_PATH, "w", encoding="utf-8") as f:
    f.write("\n".join(md_lines))
print(f"  ✓ Saved Markdown: {MD_PATH} ({os.path.getsize(MD_PATH)} bytes)")


# ==============================================================================
# 2. GENERATE MICROSOFT WORD REPORT (.docx)
# ==============================================================================
print("[2/3] Generating Microsoft Word Report (.docx)...")
doc = Document()

for section in doc.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

PRIMARY = RGBColor(15, 23, 42)
ACCENT = RGBColor(14, 116, 144)

# Title Page
p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_title = p_title.add_run(sanitize_xml(meta['title']))
run_title.font.name = 'Calibri'
run_title.font.size = Pt(24)
run_title.font.bold = True
run_title.font.color.rgb = PRIMARY

p_sub = doc.add_paragraph()
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_sub = p_sub.add_run(sanitize_xml(f"\n{meta['course']}\nAcademic Year: {meta['academic_year']}"))
run_sub.font.name = 'Calibri'
run_sub.font.size = Pt(14)
run_sub.font.color.rgb = ACCENT

p_author = doc.add_paragraph()
p_author.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_auth = p_author.add_run(sanitize_xml(f"\n\nPrepared By:\n{meta['student_name']}\nRegister Number: {meta['register_number']}\n{meta['institution']}\n{meta['date']}\n\n"))
run_auth.font.name = 'Calibri'
run_auth.font.size = Pt(12)
run_auth.font.color.rgb = PRIMARY

doc.add_page_break()

for ch in chapters:
    h = doc.add_heading(sanitize_xml(ch['title']), level=1)
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(6)
    
    lines = ch['content'].strip().split("\n")
    in_table = False
    table_rows = []

    for line in lines:
        line = line.strip()
        if not line:
            if in_table and len(table_rows) >= 2:
                headers = table_rows[0]
                rows_data = [r[:len(headers)] + [""] * max(0, len(headers) - len(r)) for r in table_rows[1:]]
                t = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                t.style = 'Light Shading Accent 1'
                for c_idx, h_text in enumerate(headers):
                    t.rows[0].cells[c_idx].text = sanitize_xml(h_text)
                    for r in t.rows[0].cells[c_idx].paragraphs[0].runs:
                        r.font.bold = True
                for r_idx, r_data in enumerate(rows_data):
                    for c_idx, val in enumerate(r_data):
                        if c_idx < len(headers):
                            t.rows[r_idx + 1].cells[c_idx].text = sanitize_xml(val)
                doc.add_paragraph()
                in_table = False
                table_rows = []
            continue

        if line.startswith("### "):
            h3_text = line.replace("### ", "").strip()
            h3 = doc.add_heading(sanitize_xml(h3_text), level=2)
            h3.paragraph_format.space_before = Pt(10)
            h3.paragraph_format.space_after = Pt(3)
        elif line.startswith("|") and "|" in line:
            if not in_table:
                in_table = True
                table_rows = []
            cells = [c.strip().replace("**", "") for c in line.split("|")[1:-1]]
            if not all(c.startswith("-") for c in cells):
                table_rows.append(cells)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(4)
            clean_text = line.replace("**", "").replace("*", "")
            p.add_run(sanitize_xml(clean_text))

    if in_table and len(table_rows) >= 2:
        headers = table_rows[0]
        rows_data = [r[:len(headers)] + [""] * max(0, len(headers) - len(r)) for r in table_rows[1:]]
        t = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.style = 'Light Shading Accent 1'
        for c_idx, h_text in enumerate(headers):
            t.rows[0].cells[c_idx].text = sanitize_xml(h_text)
            for r in t.rows[0].cells[c_idx].paragraphs[0].runs:
                r.font.bold = True
        for r_idx, r_data in enumerate(rows_data):
            for c_idx, val in enumerate(r_data):
                if c_idx < len(headers):
                    t.rows[r_idx + 1].cells[c_idx].text = sanitize_xml(val)
        doc.add_paragraph()

    def add_fig(filename, caption):
        fpath = os.path.join(FIGURES_DIR, filename)
        if os.path.exists(fpath):
            doc.add_paragraph()
            doc.add_picture(fpath, width=Inches(6.0))
            cap = doc.add_paragraph(sanitize_xml(caption))
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.runs[0].font.size = Pt(9.5)
            cap.runs[0].font.italic = True
            doc.add_paragraph()

    if ch['id'] == 4:
        add_fig("arch_system_overview.png", "Figure 4.1: QuantLens-XAI End-to-End System Architecture Overview")
    elif ch['id'] == 5:
        add_fig("arch_hybrid_database.png", "Figure 5.1: Hybrid Polyglot Persistence Architecture (SQLite 3 + MongoDB Community 7.0)")
        add_fig("07_mongodb_compass_or_cli.png", "Figure 5.2: MongoDB Community 7.0 Collection & BSON Inspection")
    elif ch['id'] == 11:
        add_fig("01_strategy_builder.png", "Figure 11.1: Strategy Builder Natural Language Prompting Interface")
        add_fig("02_decision_trace.png", "Figure 11.2: Seven-Stage Observable Decision Trace and AI Risk Advisory")
        add_fig("03_backtest_candlestick.png", "Figure 11.3: Event-Driven Backtesting Candlestick Chart, Overlays, and Trade Ledger")
        add_fig("04_xai_explainability.png", "Figure 11.4: Dual Model Predictions, Real SHAP Attributions, and Slide 9 Empirical Evaluation")
        add_fig("05_paper_broker.png", "Figure 11.5: Groww-Inspired Paper Broker Portfolio, Orders Ledger, and Audit Trail")
        add_fig("06_delta_bot_dashboard.png", "Figure 11.6: Legacy Delta Algorithmic Trading Bot Operations Dashboard")

doc.add_page_break()
doc.add_heading("Appendices", level=1)
for ap in appendices:
    doc.add_heading(sanitize_xml(ap['title']), level=2)
    clean_app = ap['content'].strip().replace("```python", "").replace("```sql", "").replace("```javascript", "").replace("```", "")
    for l in clean_app.split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.add_run(sanitize_xml(l))

doc.save(DOCX_PATH)
print(f"  ✓ Saved Word Document: {DOCX_PATH} ({os.path.getsize(DOCX_PATH)} bytes)")


# ==============================================================================
# 3. GENERATE HIGH-QUALITY PDF VIA HTML & HEADLESS CHROME
# ==============================================================================
print("[3/3] Generating High-Resolution Academic PDF (.pdf)...")

html_body = []
html_body.append(f"""
<div class="cover-page">
    <div class="academic-badge">Academic Project Report • AY {meta['academic_year']}</div>
    <h1 class="main-title">{meta['title']}</h1>
    <div class="course-subtitle">{meta['course']}</div>
    <div class="author-block">
        <p><strong>Author:</strong> {meta['student_name']}</p>
        <p><strong>Register Number:</strong> {meta['register_number']}</p>
        <p><strong>Department:</strong> {meta['institution']}</p>
        <p><strong>Date:</strong> {meta['date']}</p>
    </div>
</div>
<div class="page-break"></div>
""")

for ch in chapters:
    html_body.append(f"<div class='chapter'>")
    html_body.append(f"<h2>{ch['title']}</h2>")
    
    lines = ch['content'].strip().split("\n")
    in_table = False
    table_rows = []
    
    for line in lines:
        l = line.strip()
        if not l:
            if in_table and table_rows:
                html_body.append("<table class='data-table'>")
                for r_idx, r in enumerate(table_rows):
                    if r_idx == 0:
                        html_body.append("<tr>" + "".join(f"<th>{c}</th>" for c in r) + "</tr>")
                    else:
                        html_body.append("<tr>" + "".join(f"<td>{c}</td>" for c in r) + "</tr>")
                html_body.append("</table>")
                in_table = False
                table_rows = []
            continue
            
        if l.startswith("### "):
            html_body.append(f"<h3>{l.replace('### ', '')}</h3>")
        elif l.startswith("|") and "|" in l:
            if not in_table:
                in_table = True
                table_rows = []
            cells = [c.strip().replace("**", "") for c in l.split("|")[1:-1]]
            if not all(c.startswith("-") for c in cells):
                table_rows.append(cells)
        elif l.startswith("- "):
            html_body.append(f"<li>{l.replace('- ', '')}</li>")
        elif any(l.startswith(f"{i}. ") for i in range(1, 15)):
            html_body.append(f"<p class='numbered-item'>{l}</p>")
        else:
            html_body.append(f"<p>{l}</p>")
            
    if in_table and table_rows:
        html_body.append("<table class='data-table'>")
        for r_idx, r in enumerate(table_rows):
            if r_idx == 0:
                html_body.append("<tr>" + "".join(f"<th>{c}</th>" for c in r) + "</tr>")
            else:
                html_body.append("<tr>" + "".join(f"<td>{c}</td>" for c in r) + "</tr>")
        html_body.append("</table>")

    def get_fig_html(filename, caption):
        return f"""
        <div class="figure-container">
            <img src="figures/{filename}" alt="{caption}" />
            <div class="caption">{caption}</div>
        </div>
        """

    if ch['id'] == 4:
        html_body.append(get_fig_html("arch_system_overview.png", "Figure 4.1: QuantLens-XAI End-to-End System Architecture Overview"))
    elif ch['id'] == 5:
        html_body.append(get_fig_html("arch_hybrid_database.png", "Figure 5.1: Hybrid Polyglot Persistence Architecture (SQLite 3 + MongoDB Community 7.0)"))
        html_body.append(get_fig_html("07_mongodb_compass_or_cli.png", "Figure 5.2: MongoDB Community 7.0 Collection & BSON Inspection"))
    elif ch['id'] == 11:
        html_body.append(get_fig_html("01_strategy_builder.png", "Figure 11.1: Strategy Builder Natural Language Prompting Interface"))
        html_body.append(get_fig_html("02_decision_trace.png", "Figure 11.2: Seven-Stage Observable Decision Trace and AI Risk Advisory"))
        html_body.append(get_fig_html("03_backtest_candlestick.png", "Figure 11.3: Event-Driven Backtesting Candlestick Chart, Overlays, and Trade Ledger"))
        html_body.append(get_fig_html("04_xai_explainability.png", "Figure 11.4: Dual Model Predictions, Real SHAP Attributions, and Slide 9 Empirical Evaluation"))
        html_body.append(get_fig_html("05_paper_broker.png", "Figure 11.5: Groww-Inspired Paper Broker Portfolio, Orders Ledger, and Audit Trail"))
        html_body.append(get_fig_html("06_delta_bot_dashboard.png", "Figure 11.6: Legacy Delta Algorithmic Trading Bot Operations Dashboard"))

    html_body.append(f"</div>")

html_body.append("<div class='page-break'></div><h2>Appendices</h2>")
for ap in appendices:
    html_body.append(f"<h3>{ap['title']}</h3>")
    html_body.append(f"<pre><code>{ap['content'].strip()}</code></pre>")

full_html = f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<title>{meta['title']}</title>
<style>
    @page {{
        size: A4;
        margin: 20mm 15mm 20mm 15mm;
    }}
    body {{
        font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
        color: #1e293b;
        line-height: 1.6;
        font-size: 11pt;
        background: #fff;
    }}
    .page-break {{
        page-break-after: always;
    }}
    .cover-page {{
        min-height: 85vh;
        display: flex;
        flex-direction: column;
        justify-content: center;
        align-items: center;
        text-align: center;
        padding: 40px;
    }}
    .academic-badge {{
        background: #e0f2fe;
        color: #0369a1;
        padding: 6px 16px;
        border-radius: 20px;
        font-weight: 700;
        font-size: 12pt;
        margin-bottom: 24px;
        text-transform: uppercase;
        letter-spacing: 1px;
    }}
    .main-title {{
        font-size: 24pt;
        font-weight: 800;
        color: #0f172a;
        line-height: 1.25;
        margin-bottom: 16px;
    }}
    .course-subtitle {{
        font-size: 14pt;
        font-weight: 600;
        color: #0e7490;
        margin-bottom: 40px;
    }}
    .author-block {{
        background: #f8fafc;
        border: 1px solid #cbd5e1;
        border-radius: 12px;
        padding: 24px 36px;
        font-size: 12pt;
        text-align: left;
        width: 75%;
    }}
    .author-block p {{
        margin: 6px 0;
    }}
    h2 {{
        color: #0f172a;
        font-size: 16pt;
        border-bottom: 2px solid #0e7490;
        padding-bottom: 6px;
        margin-top: 24px;
        page-break-after: avoid;
    }}
    h3 {{
        color: #0369a1;
        font-size: 13pt;
        margin-top: 18px;
        page-break-after: avoid;
    }}
    p {{
        margin: 8px 0;
        text-align: justify;
    }}
    .numbered-item {{
        margin-left: 20px;
    }}
    .figure-container {{
        text-align: center;
        margin: 20px 0;
        page-break-inside: avoid;
    }}
    .figure-container img {{
        max-width: 95%;
        border-radius: 6px;
        border: 1px solid #cbd5e1;
        box-shadow: 0 4px 6px -1px rgba(0, 0, 0, 0.1);
    }}
    .caption {{
        font-size: 9pt;
        font-style: italic;
        color: #64748b;
        margin-top: 6px;
    }}
    .data-table {{
        width: 100%;
        border-collapse: collapse;
        margin: 16px 0;
        font-size: 10pt;
        page-break-inside: avoid;
    }}
    .data-table th {{
        background: #0f172a;
        color: #fff;
        padding: 8px 10px;
        text-align: left;
    }}
    .data-table td {{
        padding: 8px 10px;
        border-bottom: 1px solid #e2e8f0;
    }}
    .data-table tr:nth-child(even) {{
        background: #f8fafc;
    }}
    pre {{
        background: #f1f5f9;
        border: 1px solid #cbd5e1;
        border-radius: 6px;
        padding: 12px;
        font-size: 9pt;
        overflow-x: auto;
        page-break-inside: avoid;
    }}
</style>
</head>
<body>
{''.join(html_body)}
</body>
</html>
"""

with open(HTML_PATH, "w", encoding="utf-8") as f:
    f.write(full_html)
print(f"  ✓ Saved Print HTML: {HTML_PATH}")

chrome_bin = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
cmd = [
    chrome_bin,
    "--headless",
    "--disable-gpu",
    "--no-pdf-header-footer",
    f"--print-to-pdf={PDF_PATH}",
    HTML_PATH
]
subprocess.run(cmd, check=True)
print(f"  ✓ Saved Academic PDF: {PDF_PATH} ({os.path.getsize(PDF_PATH)} bytes)")

print("\n======================================================================")
print("  ALL 3 REPORT FORMATS GENERATED SUCCESSFULLY:")
print(f"  1. Markdown: {MD_PATH}")
print(f"  2. Word DOCX: {DOCX_PATH}")
print(f"  3. Vector PDF: {PDF_PATH}")
print("======================================================================\n")
